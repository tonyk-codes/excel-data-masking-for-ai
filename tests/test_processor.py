import pytest
from openpyxl import Workbook, load_workbook

from sheetveil.config import Settings
from sheetveil.processor import process_pair


def test_pair_tracks_movement(tmp_path):
    before = tmp_path / "before.xlsx"
    after = tmp_path / "after.xlsx"
    out = tmp_path / "out"
    wb = Workbook(); ws = wb.active; ws["A1"] = "Header"; ws["A2"] = "Alpha"; ws["B2"] = 12; wb.save(before)
    wb = Workbook(); ws = wb.active; ws["C4"] = "Alpha"; ws["D4"] = 12; wb.save(after)
    report = process_pair(before, after, out, Settings())
    assert report["movement_summary"]["shared_token_count"] == 2
    a = load_workbook(out / "masked_before.xlsx").active["A2"].value
    b = load_workbook(out / "masked_after.xlsx").active["C4"].value
    assert a == b and a.startswith("TEXT_")


def test_existing_outputs_require_force(tmp_path):
    before = tmp_path / "before.xlsx"
    after = tmp_path / "after.xlsx"
    out = tmp_path / "out"
    wb = Workbook(); wb.active["A1"] = "value"; wb.save(before)
    wb = Workbook(); wb.active["A1"] = "value"; wb.save(after)
    process_pair(before, after, out, Settings())
    with pytest.raises(FileExistsError, match="Use --force"):
        process_pair(before, after, out, Settings())
    process_pair(before, after, out, Settings(), force=True)
