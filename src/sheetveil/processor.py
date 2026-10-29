from __future__ import annotations

import json
import shutil
from collections import defaultdict
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .config import Settings
from .tokenizer import generate_key, load_key, save_key, token_for

SUPPORTED = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def _load(path: Path):
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {path.suffix}. Use .xlsx, .xlsm, .xltx, or .xltm.")
    return load_workbook(path, keep_vba=path.suffix.lower() in {".xlsm", ".xltm"}, data_only=False)


def _sanitize_workbook(workbook, settings: Settings) -> dict:
    removed = {"comments": 0, "hyperlinks": 0, "metadata": 0, "defined_names": 0}
    if settings.strip_metadata:
        props = workbook.properties
        for field in ("creator", "lastModifiedBy", "title", "subject", "description", "keywords", "category", "contentStatus", "identifier", "language"):
            if getattr(props, field, None):
                setattr(props, field, None)
                removed["metadata"] += 1
        now = datetime.now(timezone.utc).replace(tzinfo=None)
        props.created = now
        props.modified = now
    if settings.strip_defined_names:
        removed["defined_names"] = len(workbook.defined_names)
        workbook.defined_names.clear()
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if settings.strip_comments and cell.comment is not None:
                    cell.comment = None
                    removed["comments"] += 1
                if settings.strip_hyperlinks and cell.hyperlink is not None:
                    cell.hyperlink = None
                    removed["hyperlinks"] += 1
    return removed


def _mask_workbook(workbook, key: bytes, settings: Settings, book_label: str):
    occurrences = defaultdict(list)
    counts = {"constants": 0, "formulas_preserved": 0, "formulas_hashed": 0, "blanks": 0, "sheets": 0}
    worksheets = list(workbook.worksheets)
    for worksheet in worksheets:
        if not settings.include_hidden_sheets and worksheet.sheet_state != "visible":
            workbook.remove(worksheet)
            continue
        counts["sheets"] += 1
        for row_index, row in enumerate(worksheet.iter_rows(), start=1):
            for cell in row:
                value = cell.value
                if row_index == 1 and not settings.include_headers:
                    continue
                if value is None:
                    counts["blanks"] += 1
                    continue
                is_formula = cell.data_type == "f" or (isinstance(value, str) and value.startswith("="))
                if is_formula and settings.formula_mode == "preserve":
                    counts["formulas_preserved"] += 1
                    continue
                original = value
                token = token_for(original, key, settings.token_style, settings.token_hex_length)
                cell.value = token
                if is_formula:
                    counts["formulas_hashed"] += 1
                else:
                    counts["constants"] += 1
                occurrences[token].append({"workbook": book_label, "sheet": worksheet.title, "cell": cell.coordinate})
    return occurrences, counts


def _movement_report(before_occ, after_occ):
    before_tokens = set(before_occ)
    after_tokens = set(after_occ)
    exact = []
    for token in sorted(before_tokens & after_tokens):
        exact.append({"token": token, "before": before_occ[token], "after": after_occ[token]})
    return {
        "shared_token_count": len(exact),
        "before_only_token_count": len(before_tokens - after_tokens),
        "after_only_token_count": len(after_tokens - before_tokens),
        "shared_tokens": exact,
        "before_only_tokens": [{"token": t, "locations": before_occ[t]} for t in sorted(before_tokens - after_tokens)],
        "after_only_tokens": [{"token": t, "locations": after_occ[t]} for t in sorted(after_tokens - before_tokens)],
    }


def process_pair(before: str | Path, after: str | Path, output_dir: str | Path, settings: Settings,
                 key_file: str | Path | None = None, save_generated_key: str | Path | None = None,
                 force: bool = False) -> dict:
    before_path, after_path, out = Path(before), Path(after), Path(output_dir)
    if not before_path.is_file() or not after_path.is_file():
        raise FileNotFoundError("Both before and after workbook paths must exist.")
    if before_path.resolve() == after_path.resolve():
        raise ValueError("Before and after workbooks must be different files.")
    out.mkdir(parents=True, exist_ok=True)
    expected_outputs = [out / f"masked_before{before_path.suffix.lower()}", out / f"masked_after{after_path.suffix.lower()}"]
    if settings.write_reports:
        expected_outputs.extend(out / name for name in ("movement_map.json", "masking_report.json", "AI_INSTRUCTIONS.md"))
    existing = [path for path in expected_outputs if path.exists()]
    if existing and not force:
        names = ", ".join(path.name for path in existing)
        raise FileExistsError(f"Output already exists: {names}. Use --force to replace it.")
    if save_generated_key and Path(save_generated_key).exists() and not force:
        raise FileExistsError(f"Key already exists: {save_generated_key}. Use --force to replace it.")
    key = load_key(key_file) if key_file else generate_key()

    before_wb, after_wb = _load(before_path), _load(after_path)
    before_removed = _sanitize_workbook(before_wb, settings)
    after_removed = _sanitize_workbook(after_wb, settings)
    before_occ, before_counts = _mask_workbook(before_wb, key, settings, "before")
    after_occ, after_counts = _mask_workbook(after_wb, key, settings, "after")

    before_output = out / f"masked_before{before_path.suffix.lower()}"
    after_output = out / f"masked_after{after_path.suffix.lower()}"
    before_wb.save(before_output)
    after_wb.save(after_output)
    if save_generated_key:
        save_key(save_generated_key, key)

    movement = _movement_report(before_occ, after_occ)
    report = {
        "tool": "SheetVeil",
        "version": "1.0.0",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "key_exported": bool(save_generated_key),
        "settings": asdict(settings),
        "inputs": {"before_name": before_path.name, "after_name": after_path.name},
        "outputs": {"before": before_output.name, "after": after_output.name},
        "statistics": {"before": before_counts, "after": after_counts},
        "sanitization": {"before": before_removed, "after": after_removed},
        "movement_summary": {k: movement[k] for k in ("shared_token_count", "before_only_token_count", "after_only_token_count")},
        "warnings": [
            "Tokens reveal equality and repetition patterns by design.",
            "Preserved formulas may contain literals, paths, sheet names, or business logic.",
            "Images, charts, embedded objects, VBA, Power Query, pivot caches, and external links may contain information not represented by ordinary cells; review complex files before external sharing.",
            "HMAC tokenization supports exact-value lineage, not inference of hardcoded arithmetic or substring-derived values."
        ]
    }
    if settings.write_reports:
        (out / "movement_map.json").write_text(json.dumps(movement, indent=2), encoding="utf-8")
        (out / "masking_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
        shutil.copyfile(Path(__file__).with_name("AI_INSTRUCTIONS.template.md"), out / "AI_INSTRUCTIONS.md")
    return report
